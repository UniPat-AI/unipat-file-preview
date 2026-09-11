import sys,tempfile,json,hashlib,io
from pathlib import Path
sys.path.insert(0,'/Users/gelx/Desktop/code/unipat-file-preview/converter/src')
from converter.contracts import ConvertRequest
from converter.runtime import run_processor
root=Path(tempfile.mkdtemp(prefix='unipat-review-fixtures-'))
def run(name,data,output_limit=2*1024**3):
 p=root/name;p.write_bytes(data);out=root/(name+'-out')
 req=ConvertRequest.from_dict(dict(contract_version='1.0',job_id=name,source=dict(path=str(p),extension=p.suffix[1:],size_bytes=len(data),sha256=hashlib.sha256(data).hexdigest()),profile=dict(id='standard-v1',digest='test'),output_dir=str(out),limits=dict(output_bytes=output_limit)))
 r=run_processor(req).to_dict();return r,out
r,out=run('sample.txt',b'Actual document body');m=r['manifest'];print('TXT candidate:',m['representations'][0],json.loads((out/'text/index.json').read_text()))
r,out=run('sample.html',b'<p>before</p><input><p>after</p>');print('HTML input truncates suffix:',(out/'html/document.html').read_text())
r,out=run('broken.pdf',b'%PDF-1.4\nnot a PDF document\n%%EOF');print('Invalid PDF accepted:',r['ok'],r.get('manifest',{}).get('availability'))
from pypdf import PdfWriter
w=PdfWriter();w.add_blank_page(width=100,height=100);stream=io.BytesIO();w.write(stream);padded=stream.getvalue()+b' '*20000+b'\n%%EOF'
r,out=run('large.pdf',padded,5000);print('PDF output budget bypass:',r['ok'],len(padded),5000,r.get('error'))
from openpyxl import Workbook
wb=Workbook();ws=wb.active;ws.title='Data';ws['A1']=.25;ws['A1'].number_format='0%';ws['B1']='=1+1';ws['C1']='hidden';ws.column_dimensions['C'].hidden=True;ws.merge_cells('A2:B2');ws['A2']='merged';ws.freeze_panes='B2';f=io.BytesIO();wb.save(f)
r,out=run('sample.xlsx',f.getvalue());print('XLSX ok:',r['ok']);print('XLSX cells:',json.loads(next((out/'table').glob('*chunk*')).read_text()));print('XLSX metadata:',json.loads((out/'table/workbook.json').read_text()))
# Exactly the input name the Node orchestrator writes.
p=root/'input.bin';p.write_bytes(f.getvalue());req=ConvertRequest.from_dict(dict(contract_version='1.0',job_id='bin',source=dict(path=str(p),extension='xlsx',size_bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()),profile=dict(id='standard-v1',digest='test'),output_dir=str(root/'bin-out')));r=run_processor(req).to_dict();print('XLSX with host input.bin:',r['ok'],r.get('error'))
# Synthetic RGB DICOM, no patient metadata.
try:
 import numpy as np
 from pydicom.dataset import FileDataset,FileMetaDataset
 from pydicom.uid import ExplicitVRLittleEndian,SecondaryCaptureImageStorage,generate_uid
 meta=FileMetaDataset();meta.TransferSyntaxUID=ExplicitVRLittleEndian;meta.MediaStorageSOPClassUID=SecondaryCaptureImageStorage;meta.MediaStorageSOPInstanceUID=generate_uid()
 ds=FileDataset('rgb',{},file_meta=meta,preamble=b'\0'*128);ds.Rows=2;ds.Columns=3;ds.SamplesPerPixel=3;ds.PhotometricInterpretation='RGB';ds.PlanarConfiguration=0;ds.BitsAllocated=8;ds.BitsStored=8;ds.HighBit=7;ds.PixelRepresentation=0;ds.PixelData=np.arange(18,dtype=np.uint8).tobytes();b=io.BytesIO();ds.save_as(b,enforce_file_format=True)
 r,out=run('rgb.dcm',b.getvalue());print('RGB DICOM:',r['ok'],r.get('error'));print('RGB frames:',json.loads((out/'dicom/index.json').read_text()) if r['ok'] else '')
except ImportError as e:print('DICOM dependency unavailable:',e)
print('fixtures:',root)

from converter.processors._common import decode_bytes
original='中文内容'
d=decode_bytes(original.encode('gbk'))
print('GBK decoding:',original,repr(d.text),d.encoding,d.uncertain)
