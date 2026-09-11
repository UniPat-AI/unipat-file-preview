import io,json,hashlib,zipfile,sys
from pathlib import Path
from openpyxl import Workbook
sys.path.insert(0,'/Users/gelx/Desktop/code/unipat-file-preview/converter/src')
from converter.contracts import ConvertRequest
from converter.runtime import run_processor
from pypdf import PdfWriter
root=Path('/tmp/unipat-review2-fixtures');root.mkdir(exist_ok=True)
def convert(name,data,limit=2*1024**3):
 p=root/name;p.write_bytes(data);out=root/(name+'-out');out.mkdir(exist_ok=True)
 r=run_processor(ConvertRequest.from_dict(dict(contract_version='1.0',job_id=name,source=dict(path=str(p),extension=p.suffix[1:],size_bytes=len(data),sha256=hashlib.sha256(data).hexdigest()),profile=dict(id='standard-v1',digest='test'),output_dir=str(out),limits=dict(output_bytes=limit))))
 (root/(name+'.result.json')).write_text(json.dumps(r.to_dict(),ensure_ascii=False));return r,out
convert('real.txt',b'ACTUAL CONTENT FROM CONVERTER')
convert('real.csv',b'Name,Value\nAlice,42\n')
convert('real.ipynb',json.dumps(dict(cells=[dict(cell_type='code',source='1+1',outputs=[dict(output_type='execute_result',data={'text/plain':'2'})])])).encode())
wb=Workbook();hidden=wb.active;hidden.title='Internal';hidden.sheet_state='veryHidden';hidden['A1']='INTERNAL HIDDEN CONTENT';ws=wb.create_sheet('Business');wb.active=1
for key,value,fmt in [('A1',.25,'0.00%'),('B1',1234.5,'#,##0.00'),('C1',123,'000000'),('D1','=1+1','General'),('E1','="literal"','General')]: ws[key]=value;ws[key].number_format=fmt
ws['F1']='HIDDEN COLUMN';ws.column_dimensions['F'].hidden=True;ws['A2']='merged';ws.merge_cells('A2:B2');ws.freeze_panes='B2';out=io.BytesIO();wb.save(out)
# Add a saved numeric cache for D1: expected display is 2.
source=zipfile.ZipFile(io.BytesIO(out.getvalue()));fixed=io.BytesIO()
with zipfile.ZipFile(fixed,'w') as z:
 for name in source.namelist():
  data=source.read(name)
  if name=='xl/worksheets/sheet2.xml':data=data.replace(b'<f>1+1</f><v></v>',b'<f>1+1</f><v>2</v>')
  z.writestr(name,data)
r,out=convert('real.xlsx',fixed.getvalue());print('workbook',json.loads((out/'table/workbook.json').read_text()));print('cells',json.loads((out/'table/sheet_2-chunk-0000.json').read_text()))
w=PdfWriter();w.add_blank_page(width=100,height=100);w.add_metadata({'/Subject':'x'*20000});b=io.BytesIO();w.write(b);r,_=convert('budget.pdf',b.getvalue(),5000);print('valid PDF budget',r.ok,r.error)
w=PdfWriter();w.add_blank_page(width=100,height=100);w.encrypt('test-password');b=io.BytesIO();w.write(b);r,_=convert('encrypted.pdf',b.getvalue());print('encrypted PDF',r.ok,r.error)
